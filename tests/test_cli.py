"""End-to-end tests: build.py orchestrator as a real subprocess.

Nothing in this module hard-codes the product version: every expected
artifact name is derived from :data:`concrete_lab.__version__`, so a
version bump cannot silently break (or silently pass) these tests.
"""

import json
import subprocess
import sys
from pathlib import Path

import openpyxl

from concrete_lab import __version__
from concrete_lab.cli import ARTIFACT_PATTERN

REPO_ROOT = Path(__file__).resolve().parents[1]
BUILD = str(REPO_ROOT / "build.py")

#: Artifact names for the version under test (single source: the package).
ARTIFACT_NAME = ARTIFACT_PATTERN.format(version=__version__)
MANIFEST_NAME = Path(ARTIFACT_NAME).stem + ".json"
SHA_NAME = ARTIFACT_NAME + ".sha256"


def run_build_py(*args: str, cwd: Path = REPO_ROOT) -> subprocess.CompletedProcess:
    """Run the orchestrator exactly as a user would."""
    return subprocess.run(
        [sys.executable, BUILD, *args],
        cwd=str(cwd), capture_output=True, text=True, timeout=180,
    )


class TestBuildEndToEnd:
    def test_build_produces_verified_artifacts(self, tmp_path: Path) -> None:
        outcome = run_build_py("--output", str(tmp_path), "--no-protect")
        assert outcome.returncode == 0, outcome.stderr

        xlsx = tmp_path / ARTIFACT_NAME
        sha_file = tmp_path / SHA_NAME
        manifest = tmp_path / MANIFEST_NAME
        assert xlsx.exists() and sha_file.exists() and manifest.exists()

        # the manifest the build finalised includes the structural tier
        data = json.loads(manifest.read_text(encoding="utf-8"))
        assert data["application_version"] == __version__
        assert data["qa"]["failed"] == 0
        assert "structural" in data["qa"]["tiers"]
        assert data["standards"]["ASTM C39/C39M"] == "2026"

        # the delivered workbook is display-only (no formulas)
        workbook = openpyxl.load_workbook(str(xlsx))
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    assert not (isinstance(cell.value, str) and cell.value.startswith("="))

    def test_sha256_matches_artifact(self, tmp_path: Path) -> None:
        from concrete_lab.utils import compute_sha256

        assert run_build_py("--output", str(tmp_path), "--no-protect").returncode == 0
        xlsx = tmp_path / ARTIFACT_NAME
        recorded = (tmp_path / SHA_NAME).read_text(encoding="utf-8").split()[0]
        assert compute_sha256(xlsx) == recorded

    def test_config_flag_overrides_config_location(self, tmp_path: Path) -> None:
        """``--config`` must win over the repository-root config.yaml."""
        custom = tmp_path / "custom.yaml"
        custom.write_text(
            "project:\n  name: 'Custom Lab'\n  version: '" + __version__ + "'\n",
            encoding="utf-8",
        )
        outcome = run_build_py("--output", str(tmp_path), "--no-protect", "--config", str(custom))
        assert outcome.returncode == 0, outcome.stderr

        workbook = openpyxl.load_workbook(str(tmp_path / ARTIFACT_NAME))
        guide = workbook["00_راهنما"]
        assert guide["A1"].value == "Custom Lab"

    def test_user_input_build_records_its_data_source(self, tmp_path: Path) -> None:
        """A ``--input`` build must be traceable as user data, not demo data."""
        payload = tmp_path / "input.json"
        payload.write_text(
            json.dumps({"cases": [{"test_id": "1-2",
                                   "inputs": {"wet_mass_g": 1100, "dry_mass_g": 1000}}]}),
            encoding="utf-8",
        )
        outcome = run_build_py("--output", str(tmp_path), "--no-protect", "--input", str(payload))
        assert outcome.returncode == 0, outcome.stderr

        data = json.loads((tmp_path / MANIFEST_NAME).read_text(encoding="utf-8"))
        assert data["data_source"] == "user-input"
        assert data["data_source_path"] == payload.name

        workbook = openpyxl.load_workbook(str(tmp_path / ARTIFACT_NAME))
        moisture = workbook["03_آزمایش_1-2"]
        values = [cell.value for row in moisture.iter_rows() for cell in row]
        assert 10.0 in values  # (1100 − 1000) / 1000 × 100

    def test_input_and_demo_are_mutually_exclusive(self, tmp_path: Path) -> None:
        payload = tmp_path / "input.json"
        payload.write_text(json.dumps({"cases": []}), encoding="utf-8")
        outcome = run_build_py("--output", str(tmp_path), "--input", str(payload), "--demo")
        assert outcome.returncode != 0
        assert "--demo" in (outcome.stderr + outcome.stdout)

    def test_bad_input_file_fails_with_a_clear_message(self, tmp_path: Path) -> None:
        payload = tmp_path / "broken.json"
        payload.write_text(json.dumps({"cases": [{"test_id": "9-9", "inputs": {}}]}),
                           encoding="utf-8")
        outcome = run_build_py("--output", str(tmp_path), "--no-protect", "--input", str(payload))
        assert outcome.returncode == 1
        assert "9-9" in (outcome.stderr + outcome.stdout)


class TestValidateEndToEnd:
    def test_validate_runs_real_qa_and_passes(self) -> None:
        # restricted path keeps the nested run fast and recursion-free
        outcome = run_build_py("--validate", "tests/test_domain.py")
        assert outcome.returncode == 0, outcome.stderr
        assert "VALIDATION MODE" in outcome.stderr

    def test_validate_fails_when_a_test_fails(self, tmp_path: Path) -> None:
        bad_test = tmp_path / "test_forced_failure.py"
        bad_test.write_text("def test_broken():\n    assert False\n", encoding="utf-8")
        outcome = run_build_py("--validate", str(bad_test))
        assert outcome.returncode == 1


class TestCliSurface:
    def test_version_flag(self) -> None:
        outcome = run_build_py("--version")
        assert outcome.returncode == 0
        assert __version__ in outcome.stdout

    def test_help_documents_every_real_flag(self) -> None:
        outcome = run_build_py("--help")
        assert outcome.returncode == 0
        for flag in ("--output", "--no-protect", "--password-env", "--input", "--validate", "--config"):
            assert flag in outcome.stdout, f"{flag} is not documented in --help"
