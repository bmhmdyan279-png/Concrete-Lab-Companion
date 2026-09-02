"""Tests for the display-only Excel rendering layer."""

import pytest
from openpyxl import Workbook

from concrete_lab.config import AppConfig
from concrete_lab.domain.engine import run_demo_cases
from concrete_lab.qa.engine import QAReport
from concrete_lab.render.excel.assembler import assemble_workbook_model
from concrete_lab.render.excel.model import CellModel, RowModel, SheetModel, WorkbookModel
from concrete_lab.render.excel.protection import ProtectionManager
from concrete_lab.render.excel.renderer import ExcelRenderer
from concrete_lab.render.excel.styles import StyleManager


@pytest.fixture(scope="module")
def config() -> AppConfig:
    cfg = AppConfig.load()
    cfg.validate()
    return cfg


@pytest.fixture(scope="module")
def model(config: AppConfig) -> WorkbookModel:
    qa = {"golden": QAReport("golden-suite", passed=14, failed=0, warnings=17)}
    return assemble_workbook_model(run_demo_cases(), config, qa)


class TestAssembler:
    def test_expected_sheets_in_order(self, model: WorkbookModel) -> None:
        titles = model.sheet_titles
        assert titles[0] == "00_راهنما"
        assert titles[1] == "01_اطلاعات_آزمون"
        # seven implemented test sheets, then report/dashboard/QA/warnings/standards
        assert "17_آزمایش_4-1" in titles
        assert titles[-1] == "_Standards"
        assert "22_گزارش" in titles
        assert "24_QA_Test" in titles

    def test_sheet_titles_within_excel_limit(self, model: WorkbookModel) -> None:
        assert all(len(title) <= 31 for title in model.sheet_titles)


class TestRenderer:
    def test_rendered_workbook_has_every_sheet(self, model: WorkbookModel) -> None:
        workbook = ExcelRenderer().render(model)
        assert [ws.title for ws in workbook.worksheets] == list(model.sheet_titles)

    def test_no_formulas_anywhere(self, model: WorkbookModel) -> None:
        """Core v4 contract: the renderer displays values, never logic."""
        workbook = ExcelRenderer().render(model)
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    assert not (isinstance(value, str) and value.startswith("=")), \
                        f"formula leaked into {sheet.title}!{cell.coordinate}"

    def test_c805_sheet_shows_disclaimer_and_badge(self, model: WorkbookModel) -> None:
        workbook = ExcelRenderer().render(model)
        sheet = workbook["21_آزمایش_4-5"]
        text = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)
        assert "ESTIMATED — NOT FOR ACCEPTANCE" in text
        assert "NOT FOR ACCEPTANCE OR REJECTION" in text

    def test_rtl_and_named_styles_applied(self, model: WorkbookModel) -> None:
        workbook = ExcelRenderer().render(model)
        assert workbook["00_راهنما"].sheet_view.rightToLeft is True
        for name in ("title", "header", "pass", "warn", "fail"):
            assert name in workbook.named_styles

    def test_long_title_rejected(self) -> None:
        bad = WorkbookModel(sheets=(SheetModel(title="x" * 40),))
        with pytest.raises(ValueError):
            ExcelRenderer().render(bad)

    def test_shared_style_manager_registers_once(self) -> None:
        """Named styles are cached and re-registration is idempotent."""
        manager = StyleManager()
        workbook = Workbook()
        manager.register(workbook)
        manager.register(workbook)  # must not raise about duplicates
        assert "title" in workbook.named_styles


class TestProtection:
    def test_no_password_means_disabled(self) -> None:
        workbook = Workbook()
        assert ProtectionManager.apply(workbook, None) is False

    def test_password_protects_user_sheets_only(self) -> None:
        workbook = Workbook()
        workbook.active.title = "00_راهنما"
        workbook.create_sheet("_Standards")
        assert ProtectionManager.apply(workbook, "secret") is True
        assert workbook["00_راهنما"].protection.sheet is True
        assert workbook["_Standards"].protection.sheet is False
        assert workbook.security.lockStructure is True


class TestModels:
    def test_cell_defaults(self) -> None:
        cell = CellModel("value")
        assert cell.style == "label"
        assert cell.number_format is None

    def test_row_keeps_order(self) -> None:
        row = RowModel(cells=(CellModel("a"), CellModel("b")))
        assert [c.value for c in row.cells] == ["a", "b"]
