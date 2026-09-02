"""Style management with cached named styles.

Review feedback: the legacy code rebuilt style objects per cell.  This
manager creates each :class:`NamedStyle` exactly once and registers it
on the workbook, so styling is a batch, cache-friendly operation.
"""

from __future__ import annotations

from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side

from concrete_lab import constants


def _fill(hex_color: str) -> PatternFill:
    """Solid pattern fill helper."""
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


class StyleManager:
    """Builds and registers the workbook's named styles exactly once."""

    #: Thin grey border shared by data cells.
    BOX_BORDER: Border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )

    def __init__(self) -> None:
        self._styles: Dict[str, NamedStyle] = {}
        self._build()

    # -- Construction --------------------------------------------------------

    def _build(self) -> None:
        """Create every named style once (the style cache)."""
        fa = constants.FONT_PRIMARY
        self._add(NamedStyle(
            name="title", font=Font(name=fa, size=14, bold=True, color="1F4E79"),
            alignment=Alignment(vertical="center"),
        ))
        self._add(NamedStyle(
            name="subtitle", font=Font(name=fa, size=10, italic=True, color="595959"),
            alignment=Alignment(vertical="center", wrap_text=True),
        ))
        self._add(NamedStyle(
            name="header",
            font=Font(name=fa, size=11, bold=True, color=constants.COLOR_HEADER_FONT),
            fill=_fill(constants.COLOR_HEADER_FILL),
            border=self.BOX_BORDER,
            alignment=Alignment(horizontal="center", vertical="center"),
        ))
        self._add(NamedStyle(
            name="label", font=Font(name=fa, size=10),
            border=self.BOX_BORDER,
            alignment=Alignment(vertical="center", wrap_text=True),
        ))
        self._add(NamedStyle(
            name="value", font=Font(name=constants.FONT_NUMBERS, size=10),
            fill=_fill(constants.COLOR_VALUE_FILL),
            border=self.BOX_BORDER,
            alignment=Alignment(horizontal="center", vertical="center"),
        ))
        self._add(NamedStyle(
            name="note", font=Font(name=fa, size=9, italic=True, color="595959"),
            alignment=Alignment(vertical="center", wrap_text=True),
        ))
        self._add(NamedStyle(
            name="badge", font=Font(name=fa, size=10, bold=True, color="7F6000"),
            fill=_fill(constants.COLOR_BADGE_FILL),
            border=self.BOX_BORDER,
            alignment=Alignment(horizontal="center", vertical="center"),
        ))
        self._add(NamedStyle(
            name="pass", font=Font(name=fa, size=10, bold=True, color=constants.COLOR_PASS_FONT),
            fill=_fill(constants.COLOR_PASS_FILL),
            border=self.BOX_BORDER,
            alignment=Alignment(horizontal="center", vertical="center"),
        ))
        self._add(NamedStyle(
            name="warn", font=Font(name=fa, size=10, bold=True, color=constants.COLOR_WARN_FONT),
            fill=_fill(constants.COLOR_WARN_FILL),
            border=self.BOX_BORDER,
            alignment=Alignment(horizontal="center", vertical="center"),
        ))
        self._add(NamedStyle(
            name="fail", font=Font(name=fa, size=10, bold=True, color=constants.COLOR_FAIL_FONT),
            fill=_fill(constants.COLOR_FAIL_FILL),
            border=self.BOX_BORDER,
            alignment=Alignment(horizontal="center", vertical="center"),
        ))
        self._add(NamedStyle(
            name="pending", font=Font(name=fa, size=10, color="808080"),
            border=self.BOX_BORDER,
            alignment=Alignment(horizontal="center", vertical="center"),
        ))

    def _add(self, style: NamedStyle) -> None:
        self._styles[style.name] = style

    # -- Queries -------------------------------------------------------------

    @property
    def names(self) -> List[str]:
        """Names of all available styles."""
        return list(self._styles)

    def register(self, workbook: Workbook) -> None:
        """Register every named style on ``workbook`` (idempotent)."""
        existing = set(workbook.named_styles)
        for style in self._styles.values():
            if style.name not in existing:
                workbook.add_named_style(style)
