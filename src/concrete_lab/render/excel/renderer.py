"""The Excel renderer: writes a WorkbookModel, computes nothing.

Contract (enforced by ``qa.structural``):

* every cell receives a display value — never a formula string;
* styling goes through cached named styles;
* layout concerns (RTL, tab colour, zoom, widths) are applied here.

If a scientific rule ever needs to change, the domain engine changes;
this file stays untouched.
"""

from __future__ import annotations

from typing import Optional

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from concrete_lab import constants
from concrete_lab.render.excel.model import SheetModel, WorkbookModel
from concrete_lab.render.excel.styles import StyleManager

#: Maximum worksheet title length enforced by Excel.
EXCEL_TITLE_LIMIT: int = 31

#: Default column widths when a sheet does not override them.
DEFAULT_WIDTH_FALLBACK = {1: 36.0, 2: 18.0, 3: 12.0, 4: 14.0, 5: 48.0}


class ExcelRenderer:
    """Display-only writer from :class:`WorkbookModel` to a workbook."""

    def __init__(self, style_manager: Optional[StyleManager] = None) -> None:
        """Create the renderer with a (shared, cached) style manager."""
        self.styles = style_manager or StyleManager()

    # -- Public API ----------------------------------------------------------

    def render(self, model: WorkbookModel) -> Workbook:
        """Render the whole model into a new openpyxl workbook.

        Args:
            model: Pure-data workbook description.

        Returns:
            The rendered workbook (not yet saved).

        Raises:
            ValueError: If a sheet title exceeds Excel's 31 characters.
        """
        workbook = Workbook()
        workbook.remove(workbook.active)
        self.styles.register(workbook)

        for sheet_model in model.sheets:
            if len(sheet_model.title) > EXCEL_TITLE_LIMIT:
                raise ValueError(f"sheet title too long: {sheet_model.title!r}")
            worksheet = workbook.create_sheet(title=sheet_model.title)
            self._layout(worksheet, sheet_model)
            self._write_rows(worksheet, sheet_model)
        return workbook

    # -- Internals -----------------------------------------------------------

    def _layout(self, worksheet: Worksheet, sheet_model: SheetModel) -> None:
        """Apply layout: RTL, zoom, tab colour, column widths."""
        worksheet.sheet_view.rightToLeft = sheet_model.rtl
        worksheet.sheet_view.zoomScale = constants.DEFAULT_ZOOM
        if sheet_model.tab_color:
            worksheet.sheet_properties.tabColor = sheet_model.tab_color
        widths = dict(sheet_model.column_widths)
        for column, width in DEFAULT_WIDTH_FALLBACK.items():
            widths.setdefault(column, width)
        for column, width in widths.items():
            worksheet.column_dimensions[get_column_letter(column)].width = width

    def _write_rows(self, worksheet: Worksheet, sheet_model: SheetModel) -> None:
        """Write every row/cell value and apply its named style."""
        for row_index, row in enumerate(sheet_model.rows, start=1):
            for col_index, cell_model in enumerate(row.cells, start=1):
                cell = worksheet.cell(row=row_index, column=col_index, value=cell_model.value)
                cell.style = cell_model.style
                if cell_model.number_format:
                    cell.number_format = cell_model.number_format
